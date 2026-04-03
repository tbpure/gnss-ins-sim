from pathlib import Path

import numpy as np
import pandas as pd
from matplotlib import pyplot as plt

from gnss_ins.akf import AdaptiveLooseCouple
from gnss_ins.loose_couple import LooseCouple
from ins.run_algo import D2R
from utils.data_io import get_imu_data_from_path, get_gnss_data_from_path, get_ref_data_from_path, \
    get_att_data_from_path, pad
from utils.unit_transfer import deg2rad
from openpyxl import load_workbook


STATUS_DIMENSION = 21
IDX_DR = slice(0, 3)
IDX_DV = slice(3, 6)
IDX_PHI = slice(6, 9)
IDX_BG = slice(9, 12)
IDX_BA = slice(12, 15)
IDX_SG = slice(15, 18)
IDX_SA = slice(18, 21)


if __name__ == '__main__':
    file_path = "/Users/yangyu/PycharmProjects/gnss-ins-sim/sim_data_gen/sim_files/saved_file/default/mid-accuracy"
    save_path = "/Users/yangyu/PycharmProjects/gnss-ins-sim/gnss_ins/result/default.xlsx"
    sheet_name = 'mid'
    save_file = False
    plot_saved = True

    imu_params = {
        "G_CONST": 9.8,
        # 随机游走 (PSD)
        "ARW": 2.0e-3 * D2R / 60,  # rad/sqrt(s)
        "VRW": 0.01,  # m/s^2/sqrt(s)，根据加速度计噪声设
        # 传感器零偏
        "gyro_bias": 0.0,
        "accel_bias": 0.0,
        "sigma_bg": 0.1 * D2R / 3600,  # rad/s/sqrt(s) 零偏漂移
        "sigma_ba": 1e-3,  # m/s^2/sqrt(s)
        # 比例因子
        "gyro_scale": 1.0,
        "accel_scale": 1.0,
        "sigma_sg": 0.001,
        "sigma_sa": 0.001,
        # 一阶马尔科夫过程相关时间
        "Tgb": 100.0,
        "Tab": 100.0,
        "Tgs": 100.0,
        "Tas": 100.0,
        # 测量标准差
        "gyro_std": 0.0,
        "accel_std": 0.0,
    }
    gnss_params = {
        "pos_std": 1.0,
        "vel_std": 1.0,
    }
    imu_data = get_imu_data_from_path(file_path, ref=False)
    imu_data[:, 3:6] = np.deg2rad(imu_data[:, 3:6])
    gps_data = get_gnss_data_from_path(file_path, ref=False)
    att_data = get_att_data_from_path(file_path, ref=True)
    ref_vel = get_ref_data_from_path(file_path, ['vel'])
    ref_pos = get_gnss_data_from_path(file_path, ref = True)
    if gps_data.shape[1] != 3:
        init_vel_b = gps_data[0, 3:6]
    else:
        init_vel_b = ref_vel['vel'][0, 0:3]
    ekf = LooseCouple(imu_params, imu_data.copy(), gps_data.copy(), init_vel_b.copy(), deg2rad(att_data[0][0:3]).copy(), save=['P', 'K'])

    ekf_states = ekf.run()
    if plot_saved:
        for k in ekf.saved:
            value = ekf.saved[k]
            mat_list = np.array(value)  # shape = (T, 21, 21)
            T = mat_list.shape[0]

            plt.figure(figsize=(12, 8))

            for i in range(21):
                for j in range(21):
                    plt.plot(range(T), mat_list[:, i, j], alpha=0.5)

            plt.title(f"Element-wise Change For {k}")
            plt.xlabel("Time Step")
            plt.ylabel("Value")
            plt.grid(True)
            plt.tight_layout()
            plt.show()
    ins_result = ekf.ins.out_put
    ekf_pos = np.zeros_like(ins_result)
    for i in range(len(ekf_states)):
        s = ekf_states[i]

        # 注意：s[IDX_DR, 0] 是 EKF 对 δr (位置误差) 的估计
        # P_corrected = P_ins - δr_estimated
        ekf_pos[i] = ins_result[i] - s[IDX_DR, 0]
    ins_result = np.array(ins_result)
    plt.figure(figsize=(8, 8), dpi=120)

    akf = AdaptiveLooseCouple(imu_params, imu_data, gps_data, init_vel_b, deg2rad(att_data[0][0:3]), save=['P'])
    akf_states = akf.run()
    akf_ins = akf.ins.out_put
    akf_pos = np.zeros_like(akf_ins)
    for i in range(len(akf_states)):
        s = akf_states[i]
        akf_pos[i] = akf_ins[i] - s[IDX_DR, 0]

    plt.plot(ekf_pos[:, 1], ekf_pos[:, 0], label="EKF", linewidth=2)
    plt.plot(akf_pos[:, 1], akf_pos[:, 0], label="AKF", linewidth=2)
    # plt.plot(ins_result[:, 1], ins_result[:, 0], label="INS", linewidth=2, alpha=0.8)
    plt.plot(gps_data[:, 1], gps_data[:, 0], label="GPS", linestyle='--', linewidth=1.8, alpha=0.9)
    plt.plot(ref_pos[:, 1], ref_pos[:, 0], label="Reference", linewidth=1.5)

    plt.xlabel('Y (ECEF) [m]', fontsize=12)
    plt.ylabel('X (ECEF) [m]', fontsize=12)

    plt.title('2D Trajectory Comparison', fontsize=14)
    plt.legend(loc='best', fontsize=11)

    plt.grid(True, linestyle='--', alpha=0.5)
    plt.axis('equal')

    plt.tight_layout()
    plt.show()


    if save_file:
        padded_gps = pad(gps_data[:, 0:3], len(ins_result))
        padded_ref = pad(ref_pos[:, 0:3], len(ins_result))
        # 统一打包成一个 DataFrame（你可以按需调整列名）
        df = pd.DataFrame({
            "ekf_x": ekf_pos[:, 0],
            "ekf_y": ekf_pos[:, 1],
            "akf_x": akf_pos[:, 0],
            "akf_y": akf_pos[:, 1],
            "ins_x": ins_result[:, 0],
            "ins_y": ins_result[:, 1],
            "gps_x": padded_gps[:, 0],
            "gps_y": padded_gps[:, 1],
            "ref_x": padded_ref[:, 0],
            "ref_y": padded_ref[:, 1],
        })

        save_path = Path(save_path)

        if save_path.exists():
            # 文件已存在 → 读取已有 sheet 名
            book = load_workbook(save_path)
            existing_sheets = book.sheetnames

            # 自动避免重名
            new_name = sheet_name
            i = 1
            while new_name in existing_sheets:
                new_name = f"{sheet_name}_{i}"
                i += 1

            sheet_name = new_name

            # 追加新 sheet，不再设置 writer.book
            with pd.ExcelWriter(
                    save_path,
                    engine="openpyxl",
                    mode="a",
                    if_sheet_exists="new"
            ) as writer:
                df.to_excel(writer, sheet_name=sheet_name, index=False)

        else:
            # 文件不存在 → 创建新文件
            with pd.ExcelWriter(save_path, engine="openpyxl") as writer:
                df.to_excel(writer, sheet_name=sheet_name, index=False)

        print(f"写入完成: {save_path}, sheet = {sheet_name}")